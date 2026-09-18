"""Парсер сметы: два блока, путь l1…l5, в результат только листья."""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string
from openpyxl.worksheet.worksheet import Worksheet

from ingest.mapping import load_budget_mapping
from ingest.models import BudgetItem, ExpenseKind
from ingest.names import apply_name_aliases, match_key_for, normalize_article_name
from ingest.rejects import write_rejects
from ingest.validate import RemainWarning, apply_remain_rules, missing_required


@dataclass
class RowRec:
    row: int
    expense_kind: ExpenseKind
    segs: list[str]
    sheet_code: str
    name: str
    limit_amount: float | None
    extra: dict[str, Any]


@dataclass
class ParseResult:
    items: list[BudgetItem] = field(default_factory=list)
    skipped: list[tuple[int, str]] = field(default_factory=list)
    rejected: list[tuple[int, str]] = field(default_factory=list)
    warnings: list[RemainWarning] = field(default_factory=list)
    rejects_path: Path | None = None


def _skip_row_set(mapping: dict[str, Any]) -> tuple[set[int], set[str], set[str]]:
    rows: set[int] = set()
    headers: set[str] = set()
    totals: set[str] = set()
    for rule in mapping.get("skip_rows") or []:
        if "rows" in rule:
            rows.update(int(x) for x in rule["rows"])
        if "row" in rule:
            rows.add(int(rule["row"]))
        if rule.get("reason") == "block_header":
            headers.update(normalize_article_name(n) for n in rule.get("names") or [])
        if rule.get("reason") == "block_total":
            totals.update(normalize_article_name(n) for n in rule.get("names") or [])
    rows.update(int(x) for x in mapping["source"].get("header_rows") or [])
    return rows, headers, totals


def _segment(value: Any) -> str:
    if value is None or value == "":
        return ""
    if isinstance(value, bool):
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, int):
        return str(value)
    return str(value).strip()


def _as_float(value: Any) -> float | None | str:
    """Число, пусто (None) или 'nan' если не разобрать."""
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        return "nan"
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace(" ", "").replace(",", ".")
    if not text:
        return None
    try:
        return float(text)
    except ValueError:
        return "nan"


def _merged_map(ws: Worksheet) -> dict[tuple[int, int], tuple[int, int]]:
    out: dict[tuple[int, int], tuple[int, int]] = {}
    for rng in ws.merged_cells.ranges:
        origin = (rng.min_row, rng.min_col)
        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                out[(r, c)] = origin
    return out


def _cell(ws: Worksheet, merges: dict[tuple[int, int], tuple[int, int]], row: int, col_letter: str) -> Any:
    col = column_index_from_string(col_letter)
    origin = merges.get((row, col), (row, col))
    return ws.cell(origin[0], origin[1]).value


def _first_name(ws: Worksheet, merges: dict[tuple[int, int], tuple[int, int]], row: int, cols: list[str]) -> str:
    for col in cols:
        val = _cell(ws, merges, row, col)
        if val is None or val == "":
            continue
        text = str(val).replace("\n", " ").strip()
        if text:
            return text
    return ""


def parse_budget(
    xlsx_path: Path,
    mapping_path: Path | None = None,
    *,
    rejects_dir: Path | None = None,
    write_reject_file: bool = True,
) -> ParseResult:
    mapping = load_budget_mapping(mapping_path)
    src = mapping["source"]
    cols = mapping["columns"]
    result = ParseResult()
    skip_rows, header_names, total_names = _skip_row_set(mapping)

    wb = load_workbook(xlsx_path, data_only=True, read_only=False)
    sheet_name = src["sheet"]
    if sheet_name not in wb.sheetnames:
        raise ValueError(f"нет листа {sheet_name!r} в {xlsx_path.name}")
    ws = wb[sheet_name]
    merges = _merged_map(ws)

    kind_by_header = {
        normalize_article_name(b["header_name"]): b["expense_kind"] for b in mapping["blocks"]
    }
    current_kind: ExpenseKind | None = None
    raw_rows: list[RowRec] = []

    start = int(src["data_start_row"])
    end = int(src.get("data_end_row") or ws.max_row)
    name_cols = list(cols["name_candidates"])
    seg_cols = list(cols["code_segments"])
    limit_col = cols["amounts"]["limit_amount"]["column"]
    extra_map: dict[str, str] = dict(cols.get("extra") or {})

    for r in range(start, end + 1):
        if r in skip_rows:
            result.skipped.append((r, "skip_row"))
            continue
        name = _first_name(ws, merges, r, name_cols)
        name_norm = normalize_article_name(name) if name else ""
        if name_norm in header_names:
            current_kind = kind_by_header[name_norm]
            result.skipped.append((r, "block_header"))
            continue
        if name_norm in total_names:
            result.skipped.append((r, "block_total"))
            continue
        if not name:
            result.skipped.append((r, "no_name"))
            continue
        if name_norm.isdigit():
            result.skipped.append((r, "leftover_header"))
            continue

        segs = [_segment(_cell(ws, merges, r, c)) for c in seg_cols]
        segs = [s for s in segs if s]
        sheet_code = ".".join(segs)
        if current_kind is None:
            result.rejected.append((r, "expense_kind: unknown"))
            continue
        if not sheet_code:
            result.rejected.append((r, "нет кода статьи"))
            continue

        limit = _as_float(_cell(ws, merges, r, limit_col))
        extra: dict[str, Any] = {}
        for key, ecol in extra_map.items():
            ev = _as_float(_cell(ws, merges, r, ecol))
            if ev is None or ev == "nan":
                continue
            extra[key] = ev

        raw_rows.append(
            RowRec(
                row=r,
                expense_kind=current_kind,
                segs=segs,
                sheet_code=sheet_code,
                name=name,
                limit_amount=None if limit == "nan" else limit,
                extra=extra,
            )
        )
        if limit == "nan":
            result.rejected.append((r, "limit_amount: not_a_number"))

    by_kind: dict[ExpenseKind, list[RowRec]] = {"production": [], "management": []}
    for rec in raw_rows:
        by_kind[rec.expense_kind].append(rec)

    leaves: list[RowRec] = []
    for kind, recs in by_kind.items():
        leaves.extend(_pick_leaves(recs, mapping, result))

    by_code: dict[tuple[str, str], RowRec] = {(x.expense_kind, x.sheet_code): x for x in raw_rows}
    source_file = xlsx_path.name
    sheet = sheet_name
    year = int(src["year"])
    currency = src.get("currency", "RUB")
    amount_unit = src.get("amount_unit", "thousand_rub")

    rejected_rows = {r for r, _ in result.rejected}
    for rec in leaves:
        if rec.row in rejected_rows:
            continue
        if rec.limit_amount is None:
            result.rejected.append((rec.row, "limit_amount: empty"))
            continue
        item = _to_item(
            rec,
            by_code,
            mapping,
            source_file=source_file,
            sheet=sheet,
            year=year,
            currency=currency,
            amount_unit=amount_unit,
        )
        result.items.append(item)
    _apply_validation(result)
    if write_reject_file:
        result.rejects_path = write_rejects(
            result.rejected,
            source_file=source_file,
            sheet=sheet,
            dest_dir=rejects_dir,
        )
    return result


def _apply_validation(result: ParseResult) -> None:
    kept: list[BudgetItem] = []
    for item in result.items:
        err = missing_required(item)
        if err:
            result.rejected.append((item.row or 0, err))
            continue
        warn = apply_remain_rules(item)
        if warn:
            result.warnings.append(warn)
        kept.append(item)
    result.items = kept


def _descendants(code: str, others: list[RowRec]) -> list[RowRec]:
    pref = code + "."
    return [o for o in others if o.sheet_code.startswith(pref)]


def _pick_leaves(recs: list[RowRec], mapping: dict[str, Any], result: ParseResult) -> list[RowRec]:
    hier = mapping.get("hierarchy") or {}
    collapse = bool(hier.get("collapse_zero_stub_children", True))
    leaves: list[RowRec] = []
    stub_codes: set[str] = set()

    for rec in recs:
        desc = _descendants(rec.sheet_code, recs)
        if rec.sheet_code in stub_codes:
            continue
        if not desc:
            leaves.append(rec)
            continue
        all_zero = all((d.limit_amount or 0) == 0 for d in desc)
        if collapse and all_zero:
            leaves.append(rec)
            for d in desc:
                stub_codes.add(d.sheet_code)
                result.skipped.append((d.row, "zero_stub_child"))
            continue
        result.skipped.append((rec.row, "parent_not_leaf"))

    return [x for x in leaves if x.sheet_code not in stub_codes]


def _to_item(
    rec: RowRec,
    by_code: dict[tuple[str, str], RowRec],
    mapping: dict[str, Any],
    *,
    source_file: str,
    sheet: str,
    year: int,
    currency: str,
    amount_unit: str,
) -> BudgetItem:
    kind = rec.expense_kind
    parts = rec.sheet_code.split(".")
    ancestor_codes: list[str] = []
    path: dict[str, str | None] = {}
    parent_norm: str | None = None
    for i in range(1, len(parts) + 1):
        code = ".".join(parts[:i])
        qualified = f"{kind}:{code}"
        ancestor_codes.append(qualified)
        node = by_code.get((kind, code))
        name = node.name if node else None
        path[f"l{i}_code"] = qualified
        path[f"l{i}_name"] = name
        if i == len(parts) - 1 and node:
            parent_norm = apply_name_aliases(normalize_article_name(node.name), mapping)

    name_norm = apply_name_aliases(normalize_article_name(rec.name), mapping)
    return BudgetItem(
        source_file=source_file,
        sheet=sheet,
        row=rec.row,
        expense_kind=kind,
        sheet_code=rec.sheet_code,
        article_code=f"{kind}:{rec.sheet_code}",
        article_name=rec.name,
        article_name_norm=name_norm,
        match_key=match_key_for(name_norm, parent_norm, mapping),
        article_level=len(parts),
        year=year,
        limit_amount=rec.limit_amount or 0.0,
        amount_unit=amount_unit,
        currency=currency,
        ancestor_codes=ancestor_codes,
        extra=rec.extra,
        **{k: path.get(k) for k in (
            "l1_code", "l1_name", "l2_code", "l2_name", "l3_code", "l3_name",
            "l4_code", "l4_name", "l5_code", "l5_name",
        )},
    )
